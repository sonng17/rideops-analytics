"""Xuất báo cáo Excel từ kho dữ liệu (dw + mart).

    python scripts/export_excel.py                 # ghi reports/rideops_report.xlsx
    python scripts/export_excel.py --out bao_cao.xlsx

Đọc cấu hình kết nối DW_* từ biến môi trường / file .env (giống project 1).
Mỗi sheet = một câu hỏi phân tích; heatmap được tô màu bằng conditional formatting.
"""
from __future__ import annotations

import argparse
import os
from datetime import datetime
from pathlib import Path

import pandas as pd
import psycopg2
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from dotenv import load_dotenv

    load_dotenv()
except ImportError:
    pass

QUERIES = {
    "KPI_thang": """
        SELECT d.year AS "Năm", d.month AS "Tháng", c.company_name AS "Hãng",
               SUM(a.trip_count) AS "Số chuyến (mẫu)",
               ROUND(SUM(a.wait_seconds_sum)::numeric / NULLIF(SUM(a.wait_count), 0) / 60, 2) AS "Chờ TB (phút)",
               ROUND(SUM(a.trip_miles_sum) / SUM(a.trip_count), 2) AS "Quãng đường TB (dặm)",
               ROUND(SUM(a.driver_pay_sum) / NULLIF(SUM(a.base_fare_sum), 0), 4) AS "Tỉ lệ tài xế nhận",
               ROUND(SUM(a.driver_pay_sum) / NULLIF(SUM(a.trip_seconds_sum) / 3600.0, 0), 2) AS "Thu nhập/giờ chở khách (USD)"
        FROM dw.agg_trips_zone_hour a
        JOIN dw.dim_date d ON d.date_key = a.pickup_date_key
        JOIN dw.dim_company c ON c.company_key = a.company_key
        GROUP BY d.year, d.month, c.company_name
        ORDER BY d.year, d.month, c.company_name""",
    "Gio_x_Thu": """
        SELECT d.day_of_week, d.day_name, a.pickup_hour,
               ROUND(SUM(a.wait_seconds_sum)::numeric / NULLIF(SUM(a.wait_count), 0) / 60, 2) AS avg_wait_min
        FROM dw.agg_trips_zone_hour a
        JOIN dw.dim_date d ON d.date_key = a.pickup_date_key
        GROUP BY d.day_of_week, d.day_name, a.pickup_hour""",
    "Top_thieu_xe": """
        SELECT day_type AS "Loại ngày", pickup_hour AS "Giờ", zone_name AS "Khu vực", borough AS "Quận",
               trips_per_day AS "Chuyến/ngày (mẫu)", avg_wait_min AS "Chờ TB (phút)",
               baseline_wait_min AS "Mặt bằng TP (phút)", wait_ratio AS "Gấp mấy lần",
               p90_wait_min AS "P90 chờ (phút)", excess_wait_min_per_day AS "Phút chờ vượt/ngày",
               net_supply_per_day AS "Cung ròng/ngày"
        FROM mart.fn_top_shortage(NULL, 50)""",
    "Quan_x_Gio": """
        SELECT z.borough, a.pickup_hour,
               ROUND(SUM(a.wait_seconds_sum)::numeric / NULLIF(SUM(a.wait_count), 0) / 60, 2) AS avg_wait_min
        FROM dw.agg_trips_zone_hour a
        JOIN dw.dim_zone z ON z.zone_id = a.pu_zone_id
        GROUP BY z.borough, a.pickup_hour""",
    "Thu_nhap_tai_xe": """
        SELECT a.pickup_hour AS "Giờ",
               SUM(a.trip_count) AS "Số chuyến (mẫu)",
               ROUND(SUM(a.driver_pay_sum) / NULLIF(SUM(a.trip_seconds_sum) / 3600.0, 0), 2) AS "USD/giờ chở khách",
               ROUND(SUM(a.driver_pay_sum) / SUM(a.trip_count), 2) AS "USD/chuyến",
               ROUND(SUM(a.tips_sum) / NULLIF(SUM(a.driver_pay_sum), 0), 4) AS "Tip / thu nhập"
        FROM dw.agg_trips_zone_hour a
        GROUP BY a.pickup_hour ORDER BY a.pickup_hour""",
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
# Xanh (chờ ngắn) -> vàng -> đỏ (chờ lâu)
HEAT = ColorScaleRule(start_type="min", start_color="63BE7B",
                      mid_type="percentile", mid_value=50, mid_color="FFEB84",
                      end_type="max", end_color="F8696B")


def query(conn, sql: str) -> pd.DataFrame:
    with conn.cursor() as cur:
        cur.execute(sql)
        cols = [c.name for c in cur.description]
        df = pd.DataFrame(cur.fetchall(), columns=cols)
    # psycopg2 trả NUMERIC dạng Decimal -> đổi sang số để Excel tính được
    for col in df.columns:
        if df[col].map(type).astype(str).str.contains("Decimal").any():
            df[col] = df[col].astype(float)
    return df


def style_sheet(ws, df: pd.DataFrame, percent_cols: tuple[str, ...] = ()) -> None:
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 30)
        if col in percent_cols:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = "0.0%"


def write_heatmap(writer, df: pd.DataFrame, sheet: str, index: str, label: str, title: str,
                  index_order: list | None = None) -> None:
    pivot = df.pivot_table(index=index, columns="pickup_hour", values="avg_wait_min")
    if index_order:
        pivot = pivot.reindex(index_order)
    pivot.columns = [f"{h}h" for h in pivot.columns]
    pivot.index.name = label
    pivot.to_excel(writer, sheet_name=sheet, startrow=2)
    ws = writer.sheets[sheet]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    last_col = get_column_letter(pivot.shape[1] + 1)
    last_row = pivot.shape[0] + 3
    ws.conditional_formatting.add(f"B4:{last_col}{last_row}", HEAT)
    ws.column_dimensions["A"].width = 16
    for c in range(2, pivot.shape[1] + 2):
        ws.column_dimensions[get_column_letter(c)].width = 6
    ws.freeze_panes = "B4"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="reports/rideops_report.xlsx")
    args = parser.parse_args()
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)

    conn = psycopg2.connect(
        host=os.getenv("DW_HOST", "localhost"), port=int(os.getenv("DW_PORT", "5433")),
        dbname=os.getenv("DW_DB", "trips_dw"), user=os.getenv("DW_USER", "dw"),
        password=os.getenv("DW_PASSWORD", "dw"),
    )
    try:
        data = {name: query(conn, sql) for name, sql in QUERIES.items()}
    finally:
        conn.close()

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        info = pd.DataFrame({
            "Mục": ["Nguồn", "Phạm vi", "Lưu ý", "Tạo lúc"],
            "Nội dung": [
                "NYC TLC HVFHV, kho dw của Trip Data Pipeline",
                "Các tháng đã nạp trong dw.fact_trips",
                "Số chuyến là dữ liệu mẫu (SAMPLE_RATE); tỉ lệ và trung bình không bị ảnh hưởng",
                datetime.now().strftime("%Y-%m-%d %H:%M"),
            ],
        })
        info.to_excel(writer, sheet_name="Gioi_thieu", index=False)
        style_sheet(writer.sheets["Gioi_thieu"], info)
        writer.sheets["Gioi_thieu"].column_dimensions["B"].width = 80

        for name in ("KPI_thang", "Top_thieu_xe", "Thu_nhap_tai_xe"):
            df = data[name]
            df.to_excel(writer, sheet_name=name, index=False)
            style_sheet(writer.sheets[name], df,
                        percent_cols=("Tỉ lệ tài xế nhận", "Tip / thu nhập"))

        days = (data["Gio_x_Thu"].sort_values("day_of_week")["day_name"].drop_duplicates().tolist())
        write_heatmap(writer, data["Gio_x_Thu"], "Gio_x_Thu", "day_name", "Thứ",
                      "Thời gian chờ trung bình (phút) theo thứ x giờ đón", days)
        write_heatmap(writer, data["Quan_x_Gio"], "Quan_x_Gio", "borough", "Quận",
                      "Thời gian chờ trung bình (phút) theo quận x giờ đón")

    print(f"Đã ghi {out.resolve()}")


if __name__ == "__main__":
    main()
